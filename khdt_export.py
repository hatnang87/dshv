"""Xuất file danh sách lớp khung (Excel) từ danh sách lớp.

Ưu tiên điền vào file mẫu của TTĐT (giữ logo, định dạng, công thức Mục lục): các mẫu đều có
sheet 'Muc luc' + các sheet lớp đặt tên '1', '2', … với D7 tên lớp, D9 thời gian, D10 địa điểm,
D11 giáo viên, B8 loại hình/hình thức. Số lớp nhiều hơn số sheet của mẫu thì tự thêm sheet
(nhân bản từ sheet '1', thêm dòng tương ứng vào Mục lục). Không có mẫu thì dùng khung cơ bản
(chỉ có dữ liệu, cùng vị trí ô như mẫu) để người dùng tự copy vào mẫu.
"""
import io
import json
import math
import os
import re
import zipfile
from copy import copy

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.hyperlink import Hyperlink

from khdt_core import bo_dau, tach_ma_khoa

CAU_HINH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mau_ds_lop.json")
THU_MUC_APP = os.path.dirname(os.path.abspath(__file__))
# Streamlit Community Cloud chạy app từ /mount/src; máy chủ này ít RAM/CPU nên xuất nhẹ hơn
TREN_CLOUD = os.path.isdir("/mount/src")
MAX_WORKERS = 2 if TREN_CLOUD else 4

TEN_MUC_LUC = "Muc luc"
FONT_TEN = "Times New Roman"
FONT_CO = 12
MAU_LINK = "0563C1"


def doc_cau_hinh():
    try:
        with open(CAU_HINH, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {"thu_muc_mau": [], "mau": {}}


def tim_file_mau(ma_loai):
    """Đường dẫn file mẫu theo cấu hình (mau_ds_lop.json), hoặc None.

    'thu_muc_mau' là một đường dẫn hoặc danh sách; đường dẫn tương đối tính từ thư mục app
    (thư mục 'mau' nằm trong repo nên chạy được cả trên cloud). Thư mục đầu tiên có file khớp được dùng.
    """
    ch = doc_cau_hinh()
    tu_khoa = ch.get("mau", {}).get(ma_loai, "")
    thu_mucs = ch.get("thu_muc_mau", [])
    if isinstance(thu_mucs, str):
        thu_mucs = [thu_mucs]
    if not tu_khoa:
        return None
    for thu_muc in thu_mucs:
        if not thu_muc:
            continue
        thu_muc = thu_muc if os.path.isabs(thu_muc) else os.path.join(THU_MUC_APP, thu_muc)
        if not os.path.isdir(thu_muc):
            continue
        for ten in sorted(os.listdir(thu_muc)):
            if ten.lower().endswith(".xlsx") and not ten.startswith("~$") \
                    and tu_khoa.lower() in ten.lower():
                return os.path.join(thu_muc, ten)
    return None


# ---------------------------------------------------------------- ô / sheet lớp
def _ghi(ws, toa_do, gia_tri):
    """Ghi chuỗi vào ô (kể cả ô nằm trong vùng gộp); chuỗi bắt đầu '=' không bị hiểu là công thức."""
    cell = ws[toa_do]
    if isinstance(cell, MergedCell):
        for rng in ws.merged_cells.ranges:
            if toa_do in rng:
                cell = ws.cell(row=rng.min_row, column=rng.min_col)
                break
    text = "" if gia_tri is None else str(gia_tri)
    cell.value = text
    if text.startswith(("=", "+", "@")):
        cell.data_type = "s"


def _hinh_thuc_b8(lop):
    return f"- Loại hình/hình thức đào tạo: {lop['loai_hinh_nhan']}/{lop['hinh_thuc']}"


def _so_sheet_lop(wb):
    return sorted((n for n in wb.sheetnames if n.isdigit()), key=int)


def _dien_sheet_lop(ws, lop, ma_loai, nua_phai=True):
    _ghi(ws, "D7", lop["ten_lop"])
    _ghi(ws, "D9", lop["thoi_gian"])
    _ghi(ws, "D10", lop["dia_diem"])
    _ghi(ws, "D11", lop["giao_vien"])
    _ghi(ws, "B8", _hinh_thuc_b8(lop))
    if ma_loai == "NVM":  # khóa nhân viên mới, vd (VNBA26-PVHK08) — lấy từ cột 'Đối tượng' của KHĐT
        ma = tach_ma_khoa(lop.get("doi_tuong", ""))
        if ma:
            _ghi(ws, "L11", f"({ma})")
            if ws["K11"].has_style:  # đồng bộ định dạng với ô số lượng HV (K11) ngay cạnh
                ws["L11"]._style = copy(ws["K11"]._style)
            else:
                ws["L11"].font = Font(name=FONT_TEN, size=FONT_CO)
                ws["L11"].alignment = Alignment(vertical="center")
    chuan_hoa_sheet_lop(ws, nua_phai)


def _gop_trai(ws, vung, cot):
    """Gộp vùng một dòng (vd 'D7:F7'), căn trái; trả về tổng độ rộng các cột để tính chiều cao dòng."""
    dich = CellRange(vung)
    if vung not in {str(r) for r in ws.merged_cells.ranges}:
        for rng in list(ws.merged_cells.ranges):  # bỏ vùng gộp cũ chồng lên vùng mới
            if not rng.isdisjoint(dich):
                ws.unmerge_cells(str(rng))
        for hang in ws[vung]:
            for cell in hang[1:]:
                cell.value = None  # chỉ ô đầu giữ nội dung
        ws.merge_cells(vung)
    o_dau = ws.cell(row=dich.min_row, column=dich.min_col)
    o_dau.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return sum((ws.column_dimensions[c].width or 8.43) for c in cot)


def chuan_hoa_sheet_lop(ws, nua_phai=True):
    """Gộp D7:F7 (tên môn/khóa) và I7:O7 (bản ghi lại ở bảng điểm danh), căn trái, chiều cao dòng đủ chứa tên dài.

    nua_phai=False: bỏ qua I7:O7 (file nhanh không có bảng điểm danh bên phải).
    """
    rong = [_gop_trai(ws, "D7:F7", "DEF")]
    if nua_phai:
        rong.append(_gop_trai(ws, "I7:O7", "IJKLMNO"))
    ten = str(ws["D7"].value or "")
    dong = max(math.ceil(len(ten) * 1.15 / max(r, 1)) for r in rong + [1e9])
    if dong > 1:
        ws.row_dimensions[7].height = max(ws.row_dimensions[7].height or 15.75, 15.75 * dong)


def _chup_anh(ws):
    """Đọc dữ liệu logo của sheet một lần: Image._data() đóng luồng ảnh sau khi đọc nên phải
    gán lại luồng mới cho ảnh gốc (nếu không lúc lưu file sẽ lỗi 'closed file')."""
    anh = []
    for img in ws._images:
        data = img._data()
        img.ref = io.BytesIO(data)
        anh.append((data, copy(img.anchor), img.width, img.height))
    return anh


def _nhan_ban_sheet(wb, nguon, anh_nguon, ten):
    """Nhân bản sheet lớp (kể cả logo — openpyxl không tự copy ảnh) khi vượt số sheet của mẫu."""
    moi = wb.copy_worksheet(nguon)
    moi.title = ten
    for data, anchor, rong, cao in anh_nguon:
        ban = XlImage(io.BytesIO(data))
        ban.anchor = copy(anchor)
        ban.width, ban.height = rong, cao
        moi.add_image(ban)
    try:
        for dv in nguon.data_validations.dataValidation:
            moi.add_data_validation(copy(dv))
        for rng, rules in nguon.conditional_formatting._cf_rules.items():
            for rule in rules:
                moi.conditional_formatting.add(str(rng.sqref), copy(rule))
    except Exception:  # noqa: BLE001 — định dạng phụ, thiếu cũng không ảnh hưởng dữ liệu
        pass
    return moi


# ---------------------------------------------------------------- Mục lục
def lap_muc_luc(ws_ml, n, sao_chep_kieu=True):
    """Dòng 3.. của 'Muc luc': STT, tên lớp (công thức + hyperlink), ngày, địa điểm, giáo viên.

    Mỗi lớp i liên kết tới sheet i như các sheet có sẵn của mẫu; dòng thừa bị xóa; thống nhất font.
    """
    kieu_mau = [copy(ws_ml.cell(row=3, column=c)._style) for c in range(1, 7)]
    cuoi = max(ws_ml.max_row, 2 + n)
    for r in range(3, cuoi + 1):
        i = r - 2
        cells = [ws_ml.cell(row=r, column=c) for c in range(1, 7)]
        for cell in cells:
            if isinstance(cell, MergedCell):
                continue
            if i > n:  # dòng không dùng: xóa nội dung, link và định dạng
                cell.value = None
                cell.hyperlink = None
                cell.style = "Normal"
            elif sao_chep_kieu:
                cell._style = copy(kieu_mau[cell.column - 1])
        if i > n:
            continue
        ws_ml.cell(row=r, column=1).value = i
        for c, o in zip((2, 3, 4, 5), ("D7", "D9", "D10", "D11")):
            ws_ml.cell(row=r, column=c).value = f"='{i}'!${o[0]}${o[1:]}"
        link = ws_ml.cell(row=r, column=2)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{i}'!A1")

    # font đồng nhất toàn bộ vùng Mục lục; tiêu đề đậm; tên lớp (cột B) màu liên kết
    for r in range(1, 3 + n):
        for c in range(1, 7):
            cell = ws_ml.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            la_link = c == 2 and r >= 3
            cell.font = Font(name=FONT_TEN, size=FONT_CO, bold=r <= 2,
                             color=MAU_LINK if la_link else None,
                             underline="single" if la_link else None)


# ---------------------------------------------------------------- logo -> Mục lục
_ID_LINK = "rIdLogoMucLuc"
_RE_CNVPR = re.compile(r"<cNvPr\b([^>]*?)(/>|>)")
_RE_HLINK = re.compile(r"<a:hlinkClick\b[^>]*?(?:/>|>.*?</a:hlinkClick>)", re.S)
_LOAI_REL_LINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def gan_link_logo(xlsx_bytes, dich=f"#'{TEN_MUC_LUC}'!A1"):
    """Bấm vào logo ở mọi sheet lớp → về 'Muc luc'!A1.

    openpyxl giữ thẻ hlinkClick của logo trong mẫu nhưng ghi lại quan hệ (rId) sai — trỏ nhầm sang
    file ảnh — nên link hỏng. Ở đây ghi lại trực tiếp vào package: bỏ link cũ, gắn link mới cho
    mọi ảnh, và thêm quan hệ hyperlink vào drawing*.xml.rels (đúng cách Excel lưu).
    """
    vao = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    ra = io.BytesIO()
    with zipfile.ZipFile(ra, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in vao.infolist():
            data = vao.read(item.filename)
            if re.fullmatch(r"xl/drawings/drawing\d+\.xml", item.filename):
                x = _RE_HLINK.sub("", data.decode("utf-8"))

                def them(m):
                    thuoc_tinh, ket = m.group(1), m.group(2)
                    link = f'<a:hlinkClick r:id="{_ID_LINK}"/>'
                    return f"<cNvPr{thuoc_tinh}>{link}</cNvPr>" if ket == "/>" else f"<cNvPr{thuoc_tinh}>{link}"

                data = _RE_CNVPR.sub(them, x).encode("utf-8")
            elif re.fullmatch(r"xl/drawings/_rels/drawing\d+\.xml\.rels", item.filename):
                rel = f'<Relationship Type="{_LOAI_REL_LINK}" Target="{dich}" Id="{_ID_LINK}"/>'
                data = data.decode("utf-8").replace("</Relationships>", rel + "</Relationships>").encode("utf-8")
            zout.writestr(item, data)
    return ra.getvalue()


# ---------------------------------------------------------------- điền mẫu / khung cơ bản
def dien_mau(mau_bytes, lops, ma_loai=None):
    """Điền các lớp vào file mẫu, trả về bytes của một file xlsx.

    Vượt số sheet của mẫu thì tự thêm sheet + dòng Mục lục; thừa thì bỏ bớt.
    ma_loai: mã loại hình (NVM → điền khóa nhân viên mới vào L11).
    """
    wb = openpyxl.load_workbook(io.BytesIO(mau_bytes))
    ten_sheet = _so_sheet_lop(wb)
    if not ten_sheet or TEN_MUC_LUC not in wb.sheetnames:
        raise ValueError("File mẫu không đúng cấu trúc (cần sheet 'Muc luc' và các sheet lớp '1', '2', …).")
    n = len(lops)
    nguon = wb[ten_sheet[0]]
    if n > len(ten_sheet):  # vượt số sheet của mẫu → nhân bản thêm từ sheet '1'
        anh_nguon = _chup_anh(nguon)
        for i in range(1, n + 1):
            if str(i) not in wb.sheetnames:
                _nhan_ban_sheet(wb, nguon, anh_nguon, str(i))
    for ten in _so_sheet_lop(wb):
        if int(ten) > n:
            del wb[ten]
    for i, lop in enumerate(lops, 1):
        _dien_sheet_lop(wb[str(i)], lop, ma_loai)
    lap_muc_luc(wb[TEN_MUC_LUC], n)
    buf = io.BytesIO()
    wb.save(buf)
    return gan_link_logo(buf.getvalue())


def _viec_dien_mau(viec):
    mau_bytes, lops, ma_loai = viec
    return dien_mau(mau_bytes, lops, ma_loai)


def dien_nhieu_mau(cac_viec):
    """Chạy nhiều việc (mau_bytes, lops, ma_loai) song song — mỗi lần nạp mẫu 80 sheet + logo mất vài giây.

    Trả về list cùng thứ tự: bytes của file, hoặc Exception nếu việc đó lỗi.
    """
    def chay_tuan_tu():
        ket_qua = []
        for v in cac_viec:
            try:
                ket_qua.append(_viec_dien_mau(v))
            except Exception as e:  # noqa: BLE001 — trả lỗi về cho giao diện hiển thị
                ket_qua.append(e)
        return ket_qua

    if len(cac_viec) <= 1:
        return chay_tuan_tu()
    try:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=min(MAX_WORKERS, os.cpu_count() or 1, len(cac_viec))) as ex:
            futs = [ex.submit(_viec_dien_mau, v) for v in cac_viec]
            ket_qua = []
            for f in futs:
                try:
                    ket_qua.append(f.result())
                except Exception as e:  # noqa: BLE001
                    ket_qua.append(e)
            return ket_qua
    except Exception:  # không tạo được tiến trình con (môi trường hạn chế) → chạy lần lượt
        return chay_tuan_tu()


def tao_khung_co_ban(lops, ma_loai=None):
    """Khung chỉ có dữ liệu (không logo), cùng vị trí ô như mẫu để người dùng tự copy vào file mẫu."""
    wb = openpyxl.Workbook()
    ws_ml = wb.active
    ws_ml.title = TEN_MUC_LUC
    ws_ml["A1"] = "DANH SACH CAC SHEET"
    for c, h in enumerate(["STT", "Lớp học", "Ngày học", "Địa điểm", "Giáo viên"], 1):
        ws_ml.cell(row=2, column=c, value=h)
    for w, col in zip((6, 65, 22, 18, 28), "ABCDE"):
        ws_ml.column_dimensions[col].width = w
    for i, lop in enumerate(lops, 1):
        ws = wb.create_sheet(title=str(i))
        _ghi(ws, "B7", "- Môn/Khóa học:")
        _ghi(ws, "B9", "- Thời gian:")
        _ghi(ws, "B10", "- Địa điểm:")
        _ghi(ws, "B11", "- Giáo viên:")
        for c, h in enumerate(["STT", "Mã NV", "Họ tên", "Đơn vị", "Ghi chú"], 1):
            ws.cell(row=13, column=c, value=h).font = Font(name=FONT_TEN, size=FONT_CO, bold=True)
        for r in range(14, 29):
            ws.cell(row=r, column=1, value=r - 13)
        for w, col in zip((6, 14, 28, 14, 15, 15), "ABCDEF"):
            ws.column_dimensions[col].width = w
        _dien_sheet_lop(ws, lop, ma_loai, nua_phai=False)
        ws["A1"] = f'=HYPERLINK("#\'{TEN_MUC_LUC}\'!A1", "Trở về Mục lục")'
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.coordinate != "A1" and not cell.font.bold:
                    cell.font = Font(name=FONT_TEN, size=FONT_CO)
    lap_muc_luc(ws_ml, len(lops), sao_chep_kieu=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ten_file_an_toan(s):
    s = bo_dau(s)
    return "".join(c if c.isalnum() or c in " -_()." else "-" for c in s).strip()
