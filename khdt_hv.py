"""Tab 3: điền học viên từ file DSHV (tổng hợp nhân sự đăng ký, V.TMM-F13) vào các file danh sách lớp khung.

Python thuần (không phụ thuộc Streamlit). Luồng: doc_dshv() đọc DSHV một lần → dien_khung() cho từng file khung:
mô tả từng sheet lớp (sheet ẩn '_khop' do Tab 2 ghi, hoặc đọc D7/D9/B8/L11 nếu file cũ) → chấm điểm khớp với từng
lớp DSHV → ghi học viên, tự giãn bảng khi nhiều học viên hơn số dòng của mẫu.

DSHV mỗi sheet (Ban đầu, BSNĐ, Định kỳ, Phục hồi, BDKT, TPM, Nâng tầm…): tiêu đề ở dòng có 'KHÓA HỌC'; dòng lớp có tên ở
cột A, loại hình cột B (có thể trống), ngày bắt đầu/kết thúc cột G/H (S/C = ca sáng/chiều), cột I có thể ghi HỦY;
dòng học viên có Mã NV ở cột C, họ tên D, đội E, trung tâm F, cột G là nhãn hãng (CI/MU…) hoặc trống; cột A của dòng học
viên chỉ là ghi chú lịch các môn con của lớp.
"""
import io
import re
import zipfile
import datetime as dt
from copy import copy

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

import khdt_core
import khdt_export

NGUONG = 30           # điểm tối thiểu để coi là khớp
BIEN_DO_HOA = 3       # hai ứng viên cách nhau ≤ giá trị này coi là hòa
DIEM_TEN_CHAC = 30    # điểm tên ≥ giá trị này: khớp tên chắc; thấp hơn: chỉ chứa nhau → "gần đúng"

TT_DA_DIEN = "Đã điền"
TT_GAN_DUNG = "Đã điền (khớp gần đúng)"
TT_KHONG_CHAC = "Không chắc chắn"
TT_KHONG_CO = "Không có trong DSHV"
TT_HUY = "DSHV ghi HỦY"
TT_KHONG_BANG = "Không tìm thấy bảng học viên"
TT_KHONG_HV = "Lớp DSHV chưa có học viên"
TT_DA_DIEN_SET = (TT_DA_DIEN, TT_GAN_DUNG)

_RE_NGAY_G = re.compile(r"\d{1,2}\s*[/\-]\s*\d{1,2}")
_RE_MA_NV = re.compile(r"\d{3,6}")
_RE_HANG = re.compile(r"(?i:h[ãa]ng)\s+([A-Z0-9]{2})\b")
_RE_HANG_DCS = re.compile(r"\(([A-Z0-9]{2})\s+DCS\)")
_RE_HANG_HAU_TO = re.compile(r"[-–]\s*([A-Z0-9]{2})\s*$")
_RE_LOP_SO = re.compile(r"\blop (\d+)\b")


# ---------------------------------------------------------------- tiện ích
def norm(s):
    """Bỏ dấu, chữ thường, ký tự lạ thành khoảng trắng."""
    s = khdt_core.bo_dau(str(s or "")).lower()
    return " ".join(re.sub(r"[^a-z0-9]", " ", s).split())


def _s(v):
    """Ô Excel -> chuỗi gọn (số nguyên dạng float 17594.0 -> '17594')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).replace("\xa0", " ").split())


def _loai_nhom(text):
    """Văn bản loại hình (cột B DSHV / tên sheet / ô B8 của khung) -> BD | DK | PH | BK | None."""
    n = norm(text)
    if "phuc hoi" in n:
        return "PH"
    if "dinh ky" in n:
        return "DK"
    if "ban dau" in n or n in ("bsnd", "bsn") or "bo sung nang dinh" in n:
        return "BD"
    if "bdkt" in n or "boi duong" in n:
        return "BK"
    return None


_LOAI_LOP_KHUNG = {"NVM": "BD", "BSND": "BD", "DINH_KY": "DK", "PHUC_HOI": "PH", "BDKT": "BK"}


def _bo_ma_khoa(text):
    """Bỏ mã khóa (VNBA26-PVHL10) khỏi tên."""
    ma = khdt_core.tach_ma_khoa(text)
    if not ma:
        return text
    m = re.search(r"\s*[-–]?\s*".join(re.escape(p) for p in re.split(r"-", ma)), text, re.IGNORECASE)
    return (text[:m.start()] + " " + text[m.end():]) if m else text


def _tach_ngoac(text):
    """'Đóng mở cửa (tàu A321)' -> ('dong mo cua', {'a321'}, 'dong mo cua tau a321'):
    tên gốc bỏ mã khóa/ngoặc, tập từ trong ngoặc, và dạng phẳng (chỉ bỏ mã khóa) để so 'Quốc tế VNA' với 'Quốc tế (VNA)'."""
    text = _bo_ma_khoa(text)
    tok = set()
    for m in re.findall(r"\(([^)]*)\)", text):
        tok |= {t for t in norm(m).split() if t != "tau"}
    return norm(re.sub(r"\([^)]*\)", " ", text)), tok, norm(text)


def _doc_bytes(nguon):
    if isinstance(nguon, (bytes, bytearray)):
        return bytes(nguon)
    if hasattr(nguon, "getvalue"):
        return nguon.getvalue()
    if hasattr(nguon, "read"):
        nguon.seek(0)
        return nguon.read()
    with open(nguon, "rb") as f:
        return f.read()


def _khoang(g, h, nam_ref):
    """(đầu, cuối) từ ô ngày bắt đầu/kết thúc của DSHV ('S14/10', 'C3/11', '25/09'); năm theo nam_ref, qua năm nếu cuối < đầu."""
    d1 = khdt_core.phan_tich_ngay(g, None, nam_ref)[0]
    d2 = khdt_core.phan_tich_ngay(h, None, nam_ref)[0]
    if not d1 and not d2:
        return None
    dau, cuoi = min(d1 or d2), max(d2 or d1)
    if cuoi < dau and (dau - cuoi).days > 150:
        try:
            cuoi = cuoi.replace(year=cuoi.year + 1)
        except ValueError:
            pass
    return (dau, cuoi) if cuoi >= dau else (cuoi, dau)


# ---------------------------------------------------------------- đọc DSHV
def doc_dshv(nguon, nam_ref=None):
    """Đọc DSHV → list lớp: {sheet, dong, ten[], ma_khoa, loai, ca, khoang, huy, hang, mon_con(dict tên→[khoảng ngày]), hv[]}.

    hv: {manv, hoten, doi, trungtam, nhan}. nguon: bytes / file upload / đường dẫn.
    """
    nam_ref = nam_ref or dt.date.today().year
    wb = openpyxl.load_workbook(io.BytesIO(_doc_bytes(nguon)), data_only=True, read_only=True)
    lops = []
    for ws in wb.worksheets:
        if norm(ws.title).replace(" ", "") in ("mucluc", "sheet1"):
            continue
        rows = [list(r) + [None] * 10 for r in ws.iter_rows(values_only=True)]
        dau = next((i for i, r in enumerate(rows) if any("khoa hoc" in norm(c) for c in r[:3] if c is not None)), None)
        if dau is None:
            continue
        loai_sheet = _loai_nhom(ws.title)
        loai_truoc = None
        cur = None
        for so_dong, r in enumerate(rows[dau + 1:], start=dau + 2):
            a, b, c, d, e, f, g, h, i = r[:9]
            ma = _s(c)
            if _RE_MA_NV.fullmatch(ma):
                if cur is not None and _s(d):
                    cur["hv"].append(dict(manv=ma, hoten=_s(d), doi=_s(e), trungtam=_s(f), nhan=_s(g)))
                if cur is not None and _s(a):
                    _ghi_chu_mon_con(cur, a, nam_ref)
                continue
            ten_a = [x.strip() for x in str(a or "").replace("\xa0", " ").split("\n") if x.strip()]
            if ten_a and g is not None and _RE_NGAY_G.search(str(g)):
                loai = _loai_nhom(b) or loai_truoc or loai_sheet
                loai_truoc = _loai_nhom(b) or loai_truoc
                ten = list(ten_a) + ([" ".join(ten_a)] if len(ten_a) > 1 else [])
                hang = None
                if "hang" in norm(ten_a[0]).split():
                    m = _RE_HANG_HAU_TO.search(ten_a[0])
                    hang = m.group(1) if m else None
                if hang:
                    ten = [_RE_HANG_HAU_TO.sub("", t).strip() for t in ten]
                ca = _s(g)[:1].upper()
                cur = dict(sheet=ws.title, dong=so_dong, ten=ten, ma_khoa=khdt_core.tach_ma_khoa(" ".join(ten_a)),
                           loai=loai, ca=ca if ca in ("S", "C") else "", khoang=_khoang(g, h, nam_ref),
                           huy="huy" in norm(i), hang=hang, mon_con={}, hv=[])
                lops.append(cur)
            elif ten_a:
                nd = norm(a)
                if nd.startswith(("nguoi lap", "ho va ten", "ngay")):
                    cur = None
                elif cur is not None:
                    _ghi_chu_mon_con(cur, a, nam_ref)
    return lops


def _ghi_chu_mon_con(cur, a, nam_ref):
    """Ghi chú cột A ('Thực hành+ Kiểm tra: 02/10 -27/10') -> tên môn con (chuẩn hóa) và khoảng ngày của lớp."""
    for dong in str(a).split("\n"):
        ten, _, ngay = dong.partition(":")
        ten = norm(ten)
        if ten and not ten.startswith(("gv", "dia diem")):
            ds = cur["mon_con"].setdefault(ten, [])
            kh = khdt_core.khoang_ngay(ngay, nam_ref=nam_ref) if ngay.strip() else None
            if kh:
                ds.append(kh)


# ---------------------------------------------------------------- mô tả sheet lớp của file khung
def _doc_meta(wb):
    if khdt_export.TEN_SHEET_META not in wb.sheetnames:
        return {}
    ws = wb[khdt_export.TEN_SHEET_META]
    meta = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] is not None:
            meta[str(r[0])] = dict(zip(khdt_export.COT_META, [("" if v is None else str(v)) for v in r]))
    return meta


def _ngay_iso(s):
    try:
        return dt.date.fromisoformat(s)
    except (TypeError, ValueError):
        return None


def mo_ta_sheet(ws, meta=None):
    """Thông tin để khớp của một sheet lớp: ten_lop, mon, chas[], ma_khoa, loai, khoang, ca, hang, doi_tuong, co_meta."""
    ten_lop = _s(ws["D7"].value)
    if not ten_lop:
        return None
    parts = [p.strip() for p in ten_lop.split("/") if p.strip()]
    mon = parts[0] if parts else ten_lop
    chas = list(parts[1:])
    if meta:
        if meta.get("muc_cha"):
            chas.append(meta["muc_cha"])
        d1, d2 = _ngay_iso(meta.get("tu_ngay")), _ngay_iso(meta.get("den_ngay"))
        khoang = (d1, d2) if d1 and d2 else None
        loai, ca, ma_khoa, doi_tuong = _LOAI_LOP_KHUNG.get(meta.get("loai_hinh")), meta.get("ca", ""), \
            meta.get("ma_khoa", ""), meta.get("doi_tuong", "")
    else:
        d9 = _s(ws["D9"].value)
        nam = re.search(r"/(\d{4})", d9)
        khoang = khdt_core.khoang_ngay(d9, nam_ref=int(nam.group(1)) if nam else None)
        ca = d9[:1].upper() if d9[:1].upper() in ("S", "C") else ""
        b8 = _s(ws["B8"].value)
        loai = _loai_nhom(b8.split(":", 1)[-1].rsplit("/", 1)[0])
        ma_khoa = khdt_core.tach_ma_khoa(_s(ws["L11"].value))
        doi_tuong = ""
    m = _RE_HANG.search(ten_lop) or _RE_HANG_DCS.search(ten_lop)
    return dict(ten_lop=ten_lop, mon=mon, chas=chas, ma_khoa=ma_khoa, loai=loai, khoang=khoang, ca=ca,
                hang=m.group(1) if m else None, doi_tuong=doi_tuong, co_meta=bool(meta))


def ke_thua_hang(mo_tas):
    """Lớp không ghi hãng nhưng cùng mục cha với các lớp có hãng → lấy hãng của lớp kết thúc ngay trước ngày bắt đầu
    (vd 'Hệ thống… (JL DCS)' 3–6/10 → 'Thực hành và kiểm tra' 7–22/10 là của JL)."""
    nhom = {}
    for mt in mo_tas:
        if mt and mt["chas"]:
            nhom.setdefault(norm(mt["chas"][-1]), []).append(mt)
    for ds in nhom.values():
        co_hang = [x for x in ds if x["hang"] and x["khoang"]]
        if not co_hang:
            continue
        for mt in ds:
            if mt["hang"] or not mt["khoang"]:
                continue
            hom_truoc = mt["khoang"][0] - dt.timedelta(days=1)
            # lớp kết thúc ngay trước đó; bỏ lớp đã có lớp kế tiếp cùng hãng (vd 'Quy định chính sách hãng NH' 13–14/10
            # đã có 'Hệ thống… (NH DCS)' bắt đầu 15/10 thì không phải lớp đứng ngay trước 'Thực hành' của hãng khác)
            ung_vien = {x["hang"] for x in co_hang if x["khoang"][1] == hom_truoc and not any(
                y is not x and y["hang"] == x["hang"] and y["khoang"][0] == hom_truoc + dt.timedelta(days=1)
                for y in co_hang)}
            if len(ung_vien) == 1:
                mt["hang"] = ung_vien.pop()


# ---------------------------------------------------------------- khớp lớp
def _giao(a, b):
    return a[0] <= b[1] and b[0] <= a[1]


def _trong(a, b):
    return b[0] <= a[0] and a[1] <= b[1]


def _lop_so(text):
    return set(_RE_LOP_SO.findall(norm(text)))


def diem_khop(mt, d):
    """(điểm, điểm_tên, lý_do) giữa mô tả sheet khung `mt` và lớp DSHV `d`; None nếu loại."""
    ly_do, diem = [], 0
    mon_b, mon_t, mon_f = _tach_ngoac(mt["mon"])
    cha_bt = [_tach_ngoac(c) for c in mt["chas"]]
    day_du = norm(mt["ten_lop"])
    khung, dshv = mt["khoang"], d["khoang"]

    # --- tên
    tot = 0
    for raw in d["ten"]:
        b_d, t_d, f_d = _tach_ngoac(raw)
        if not b_d:
            continue
        if f_d == day_du or f_d == mon_f:
            s, lydo = 40, "trùng tên"
        elif b_d == mon_b and mon_t <= t_d:
            s, lydo = (40, "trùng tên") if mon_t == t_d else (34, "trùng tên (biến thể trong ngoặc)")
        elif any(f_d == cf or (b_d == cb and ct <= t_d) for cb, ct, cf in cha_bt if cb):
            s, lydo = 30, "trùng tên khóa/mục cha"
        elif len(b_d) > 12 and (f" {b_d} " in f" {mon_b} " or f" {mon_b} " in f" {b_d} "):
            s, lydo = 12, "tên chứa nhau"
        elif len(b_d) > 12 and any(cb and (f" {b_d} " in f" {cb} " or f" {cb} " in f" {b_d} ") for cb, _, _ in cha_bt):
            s, lydo = 12, "mục cha chứa nhau"
        else:
            continue
        if s > tot:
            tot, ten_ly_do = s, lydo
    # tên môn chỉ có trong ghi chú lịch lớp: chỉ dùng khi khung không có khóa/mục cha thật (mục cha viết hoa toàn bộ như
    # 'CHƯƠNG TRÌNH ĐÀO TẠO BỔ SUNG' là tiêu đề nhóm), tránh gán nhầm các lớp tên chung 'Thực hành + kiểm tra'
    co_cha_that = any(c and not c.isupper() for c in mt["chas"])
    if tot < 25 and not co_cha_that and norm(mt["mon"]) in d["mon_con"] and khung and dshv and _trong(khung, dshv):
        tot, ten_ly_do = 25, "môn con trong ghi chú lịch lớp"
    lich = 0  # ghi chú lịch môn con của lớp DSHV ('Thực hành+ Kiểm tra: 02/10 -27/10') trùng/chứa ngày của khung
    if tot and khung:
        for ten_con in [norm(mt["mon"])] + [norm(c) for c in mt["chas"]]:
            for kh in d["mon_con"].get(ten_con, []):
                lich = max(lich, 20 if kh == khung else 10 if _trong(khung, kh) else 0)
    khop_ma = bool(mt["ma_khoa"] and d["ma_khoa"] and mt["ma_khoa"] == d["ma_khoa"])
    if khop_ma:
        diem += 60
        ly_do.append(f"mã khóa {d['ma_khoa']}")
    elif tot == 0:
        return None
    if tot:
        diem += tot + lich
        ly_do.append(ten_ly_do + (" + khớp lịch môn con" if lich else ""))

    # --- "Lớp N" khác nhau thì không thể là một
    so_khung, so_dshv = _lop_so(mt["mon"]), set().union(*[_lop_so(t) for t in d["ten"]])
    if so_khung and so_dshv and not (so_khung & so_dshv):
        return None
    if so_dshv and not so_khung:
        diem -= 6

    # --- ngày
    if khung and dshv:
        if _trong(khung, dshv):
            diem += 10
            ly_do.append("ngày nằm trong khoảng lớp")
        elif _giao(khung, dshv):
            diem += 3
            ly_do.append("ngày giao nhau")
        elif not khop_ma:
            return None
        if khung == dshv:
            diem += 6
    mot_ngay = khung and dshv and khung[0] == khung[1] and dshv[0] == dshv[1]  # S/C chỉ có nghĩa với lớp 1 ngày
    if mot_ngay and mt["ca"] and d["ca"] and mt["ca"] != d["ca"]:
        diem -= 20
        ly_do.append("khác ca sáng/chiều")

    # --- loại hình (mềm: cột B/tên sheet của DSHV không nhất quán)
    if mt["loai"] and d["loai"]:
        diem += 8 if mt["loai"] == d["loai"] else -8

    # --- hãng
    if mt["hang"]:
        if d["hang"] == mt["hang"]:
            diem += 30 if (khung and dshv and _trong(khung, dshv)) else 5
            ly_do.append(f"hãng {mt['hang']}")
        elif d["hang"]:
            diem -= 40
    elif d["hang"]:
        diem -= 15

    # --- đối tượng của khung (PVSĐ (DVSĐ)…) khớp trung tâm/đội của học viên: phân biệt lớp trùng tên
    if mt["doi_tuong"] and d["hv"]:
        tu = set(norm(mt["doi_tuong"]).split())
        dung = sum(1 for h in d["hv"] if norm(h["trungtam"]) in tu or norm(h["doi"]) in tu)
        diem += 6 * dung / len(d["hv"])
    return diem, tot, ly_do


def khop_lop(mt, ds_dshv):
    """Chọn lớp DSHV cho một sheet khung → dict(trang_thai, lop, diem, ly_do, ung_vien[], hv[])."""
    xep = []
    for d in ds_dshv:
        kq = diem_khop(mt, d)
        if kq and kq[0] >= NGUONG:
            xep.append((kq[0], kq[1], kq[2], d))
    if not xep:
        return dict(trang_thai=TT_KHONG_CO, lop=None, diem=0, ly_do="", ung_vien=[], hv=[])
    xep.sort(key=lambda x: -x[0])
    diem, tot, ly_do, d = xep[0]
    hoa = [x for x in xep if x[0] >= diem - BIEN_DO_HOA]
    ung_vien = [x[3] for x in hoa]
    if len({tuple(h["manv"] for h in x[3]["hv"]) for x in hoa}) > 1 or any(x[3]["huy"] != d["huy"] for x in hoa):
        return dict(trang_thai=TT_KHONG_CHAC, lop=None, diem=diem, ung_vien=ung_vien,
                    ly_do=f"{len(hoa)} lớp DSHV điểm tương đương: " + "; ".join(
                        f"{x[3]['sheet']}!{x[3]['dong']} {x[3]['ten'][0][:40]}" for x in hoa), hv=[])
    if d["huy"]:
        return dict(trang_thai=TT_HUY, lop=d, diem=diem, ly_do="; ".join(ly_do), ung_vien=[d], hv=[])
    hv = d["hv"]
    if mt["hang"] and not d["hang"] and any(h["nhan"] for h in hv):
        hv = [h for h in hv if h["nhan"].upper() == mt["hang"]]
        ly_do = ly_do + [f"lọc học viên nhãn hãng {mt['hang']}"]
    if not hv:
        return dict(trang_thai=TT_KHONG_HV, lop=d, diem=diem, ly_do="; ".join(ly_do), ung_vien=[d], hv=[])
    return dict(trang_thai=TT_DA_DIEN if tot >= DIEM_TEN_CHAC or "mã khóa" in " ".join(ly_do) else TT_GAN_DUNG,
                lop=d, diem=diem, ly_do="; ".join(ly_do), ung_vien=[d], hv=hv)


# ---------------------------------------------------------------- điền vào sheet lớp
def tim_bang_hoc_vien(ws):
    """(dòng đầu dữ liệu, cột 'Mã NV') hoặc None. Mẫu TTĐT: tiêu đề gộp 2 dòng (12-13), dữ liệu từ dòng 14, STT ở cột B."""
    for r in range(8, 16):
        for c in range(1, 8):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and norm(v).replace(" ", "") == "manv":
                bat_dau = r + 1
                while isinstance(ws.cell(row=bat_dau, column=c), MergedCell):
                    bat_dau += 1
                return bat_dau, c
    return None


def _dong_cuoi_bang(ws, bat_dau, c_stt):
    """Dòng cuối của bảng = dòng cuối liên tiếp có STT (công thức =ROW()-13 hoặc số)."""
    r = bat_dau
    while r <= ws.max_row:
        v = ws.cell(row=r, column=c_stt).value
        if isinstance(v, (int, float)) or (isinstance(v, str) and v.upper().startswith("=ROW")):
            r += 1
        else:
            break
    return r - 1 if r > bat_dau else bat_dau + 14


def _mo_rong_bang(ws, cuoi, k):
    """Chèn k dòng vào cuối bảng học viên: dời khối chân trang (chữ ký, lưu ý) xuống, nhân style/công thức của dòng cuối."""
    chan = cuoi + 1
    max_row, max_col = ws.max_row, ws.max_column
    chieu_cao = {r: ws.row_dimensions[r].height for r in range(chan, max_row + 1)}
    goc = {c: (ws.cell(row=cuoi, column=c)._style, ws.cell(row=cuoi, column=c).value) for c in range(1, max_col + 1)}
    cao_cuoi = ws.row_dimensions[cuoi].height
    if max_row >= chan:
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row >= chan:
                rng.shift(row_shift=k)
        ws.move_range(f"A{chan}:{get_column_letter(max_col)}{max_row}", rows=k, cols=0)
        for r in range(max_row, chan - 1, -1):
            ws.row_dimensions[r + k].height = chieu_cao[r]
    for r in range(chan, chan + k):
        ws.row_dimensions[r].height = cao_cuoi
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            style, val = goc[c]
            cell._style = copy(style)
            if isinstance(val, str) and val.startswith("="):
                col = get_column_letter(c)
                cell.value = Translator(val, origin=f"{col}{cuoi}").translate_formula(f"{col}{r}")
    # công thức đếm sĩ số ở đầu sheet ($L$14:$L47) phải bao tới dòng cuối mới
    moi = cuoi + k
    for hang in ws.iter_rows(min_row=1, max_row=chan - 1):
        for cell in hang:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and "COUNT" in v.upper():
                cell.value = re.sub(r"(\$L\$\d+:\$L)(\d+)", lambda m: m.group(1) + str(max(int(m.group(2)), moi)), v)


def _ghi_hoc_vien(ws, vi_tri, hv):
    bat_dau, c_ma = vi_tri
    cuoi = _dong_cuoi_bang(ws, bat_dau, c_ma - 1)
    thieu = len(hv) - (cuoi - bat_dau + 1)
    if thieu > 0:
        _mo_rong_bang(ws, cuoi, thieu)
        cuoi += thieu
    for r in range(bat_dau, cuoi + 1):
        for c in range(c_ma, c_ma + 3):
            cell = ws.cell(row=r, column=c)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for i, h in enumerate(hv):
        r = bat_dau + i
        if ws.cell(row=r, column=c_ma - 1).value is None:  # mẫu đã có công thức STT thì giữ
            ws.cell(row=r, column=c_ma - 1, value=i + 1)
        ws.cell(row=r, column=c_ma, value=h["manv"])
        ws.cell(row=r, column=c_ma + 1, value=h["hoten"])
        ws.cell(row=r, column=c_ma + 2, value=h["trungtam"])


def dien_khung(du_lieu, ds_dshv, ten_file=""):
    """Điền học viên vào mọi sheet lớp của một file khung → (bytes file mới, thống kê theo sheet).

    thống kê: {file, sheet, ten_lop, trang_thai, so_hv, lop_dshv, diem, ly_do, khoa_dshv[], khoa_ung_vien[]}.
    """
    wb = openpyxl.load_workbook(io.BytesIO(_doc_bytes(du_lieu)))
    meta = _doc_meta(wb)
    cac_sheet = [ws for ws in wb.worksheets if ws.title not in (khdt_export.TEN_MUC_LUC, khdt_export.TEN_SHEET_META)
                 and norm(ws.title).replace(" ", "") != "mucluc"]
    mo_tas = [mo_ta_sheet(ws, meta.get(ws.title)) for ws in cac_sheet]
    ke_thua_hang(mo_tas)
    thong_ke = []
    for ws, mt in zip(cac_sheet, mo_tas):
        if not mt:
            continue
        kq = khop_lop(mt, ds_dshv)
        trang_thai, so_hv = kq["trang_thai"], 0
        if trang_thai in TT_DA_DIEN_SET:
            vi_tri = tim_bang_hoc_vien(ws)
            if vi_tri:
                _ghi_hoc_vien(ws, vi_tri, kq["hv"])
                so_hv = len(kq["hv"])
            else:
                trang_thai = TT_KHONG_BANG
        d = kq["lop"]
        thong_ke.append(dict(
            file=ten_file, sheet=ws.title, ten_lop=mt["ten_lop"], trang_thai=trang_thai, so_hv=so_hv,
            lop_dshv=f"{d['sheet']}!{d['dong']}: {d['ten'][0][:60]}" if d else "", diem=round(kq["diem"], 1),
            ly_do=kq["ly_do"], co_meta=mt["co_meta"],
            khoa_dshv=[(d["sheet"], d["dong"])] if d and trang_thai in TT_DA_DIEN_SET else [],
            khoa_ung_vien=[(x["sheet"], x["dong"]) for x in kq["ung_vien"]]))
    buf = io.BytesIO()
    wb.save(buf)
    try:  # openpyxl ghi lại quan hệ link của logo sai → gắn lại link về 'Muc luc'
        ket_qua = khdt_export.gan_link_logo(buf.getvalue())
    except Exception:  # noqa: BLE001 — chỉ là link logo
        ket_qua = buf.getvalue()
    return ket_qua, thong_ke


def lop_dshv_chua_dung(ds_dshv, thong_ke):
    """Lớp DSHV có học viên (không HỦY) chưa gắn với sheet khung nào — kể cả ứng viên của ca 'không chắc chắn'."""
    dung = {k for t in thong_ke for k in t["khoa_dshv"] + t["khoa_ung_vien"]}
    return [d for d in ds_dshv if d["hv"] and not d["huy"] and (d["sheet"], d["dong"]) not in dung]


# ---------------------------------------------------------------- gom file khung
def gom_file_khung(tu_tab2, uploads):
    """list (tên, bytes): file từ Tab 2 + file upload (.xlsx hoặc .zip); trùng tên thì ưu tiên bản upload."""
    ds = {ten: data for ten, data in tu_tab2}
    for ten, data in uploads:
        if ten.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for it in zf.infolist():
                    base = it.filename.replace("\\", "/").split("/")[-1]
                    if it.is_dir() or "__MACOSX" in it.filename or base.startswith(("~$", ".")) \
                            or not base.lower().endswith(".xlsx"):
                        continue
                    ds[base] = zf.read(it)
        elif ten.lower().endswith(".xlsx") and not ten.startswith("~$"):
            ds[ten] = data
    return list(ds.items())
