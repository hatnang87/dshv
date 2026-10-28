"""Kiểm thử nhanh khdt_hv với dữ liệu giả lập: `python test_khdt_hv.py` (hoặc pytest)."""
import io
import os
import zipfile

import openpyxl

import khdt_export
import khdt_hv


def _dshv_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ban đầu"
    ws.append([None, "Biểu mẫu"])
    ws.append(["KHÓA HỌC", "LOẠI HÌNH ĐT", "MÃ NHÂN VIÊN", "HỌ VÀ TÊN", "ĐỘI", "TRUNG TÂM", "NGÀY BẮT ĐẦU", "NGÀY KẾT THÚC"])
    ws.append(["Nghiệp vụ A -VNBA26-ABC10", "Ban đầu", None, None, None, None, "25/09", "14/10"])
    ws.append([None, None, "MãNV", None, None, None, "Mã LV", None])
    for i in range(20):  # 20 học viên > 15 dòng của file nhanh
        ghi_chu = "Thực hành+ Kiểm tra: 02/10 -10/10" if i == 0 else None
        ws.append([ghi_chu, None, str(17000 + i), f"Học viên {i}", "QT1", "PVHK", None, None])
    ws.append(["Lớp bị hủy", "Định kỳ", None, None, None, None, "S7/10", "S7/10", "HỦY"])
    ws.append([None, None, "18000", "Người hủy", "QT1", "PVHK", None, None])
    ws.append(["Người lập", "Phòng KHHC"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _lop(ten, tu, den, thoi_gian, doi_tuong="", loai="NVM", muc="A > B > Nghiệp vụ A"):
    return dict(ten_lop=ten, thoi_gian=thoi_gian, dia_diem="TTĐT", giao_vien="GV", loai_hinh_nhan="Ban đầu", hinh_thuc="Trực tiếp",
                doi_tuong=doi_tuong, loai_hinh=loai, tu_ngay=tu, den_ngay=den, muc=muc)


def test_doc_dshv():
    ds = khdt_hv.doc_dshv(_dshv_bytes(), 2026)
    assert len(ds) == 2  # ghi chú ở cột A của dòng học viên không tạo lớp; "Người lập" bỏ qua
    assert ds[0]["ma_khoa"] == "VNBA26-ABC10" and len(ds[0]["hv"]) == 20 and not ds[0]["huy"]
    assert "thuc hanh kiem tra" in ds[0]["mon_con"]
    assert ds[1]["huy"] and ds[1]["ca"] == "S"


def test_dien_khung_nhanh_gian_bang():
    ds = khdt_hv.doc_dshv(_dshv_bytes(), 2026)
    lops = [_lop("Môn 1/Nghiệp vụ A", "2026-10-01", "2026-10-02", "1-2/10/2026", "VNBA26- ABC10"),   # khớp theo mã khóa
            _lop("Lớp bị hủy", "2026-10-07", "2026-10-07", "S7/10/2026", loai="DINH_KY"),
            _lop("Không có trong DSHV", "2026-10-03", "2026-10-03", "3/10/2026", muc="A > B > Khóa khác")]
    ket_qua, tk = khdt_hv.dien_khung(khdt_export.tao_khung_co_ban(lops, "NVM"), ds, "x.xlsx")
    assert [t["trang_thai"] for t in tk] == [khdt_hv.TT_DA_DIEN, khdt_hv.TT_HUY, khdt_hv.TT_KHONG_CO]
    assert tk[0]["so_hv"] == 20
    ws = openpyxl.load_workbook(io.BytesIO(ket_qua))["1"]
    assert ws["B14"].value == "17000" and ws["B33"].value == "17019" and ws["D33"].value == "PVHK"
    assert ws["A33"].value == 20  # bảng giãn từ 15 lên 20 dòng, STT liên tục
    assert khdt_hv.lop_dshv_chua_dung(ds, tk) == []


def _dshv_nho_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Định kỳ"
    ws.append(["KHÓA HỌC", "LOẠI HÌNH ĐT", "MÃ NHÂN VIÊN", "HỌ VÀ TÊN", "ĐỘI", "TRUNG TÂM", "NGÀY BẮT ĐẦU", "NGÀY KẾT THÚC"])
    ws.append(["Môn nhỏ", "Định kỳ", None, None, None, None, "5/10", "5/10"])
    for i in range(5):
        ws.append([None, None, str(18000 + i), f"HV nhỏ {i}", "QT1", "PVHK", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_thu_gon_bang_so_hv_dung_chung():
    """Mẫu 20 dòng mà lớp chỉ 5 học viên: bảng co về 5 dòng, chân trang liền sau; Mục lục có Số HV; 2 sheet dùng chung 1 lớp DSHV."""
    mau = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mau", "DS lOP HOC ĐINH KỲ  (T 2026).xlsx")
    with open(mau, "rb") as f:
        lops = [_lop("Môn nhỏ", "2026-10-05", "2026-10-05", "5/10/2026", loai="DINH_KY") for _ in range(2)]
        lops.append(_lop("Môn không có", "2026-10-09", "2026-10-09", "9/10/2026", loai="DINH_KY", muc="A > B > Khóa khác"))
        khung = khdt_export.dien_mau(f.read(), lops, "DINH_KY")
    ds = khdt_hv.doc_dshv(_dshv_nho_bytes(), 2026)
    ket_qua, tk = khdt_hv.dien_khung(khung, ds, "k.xlsx")
    khdt_hv.gan_dung_chung(tk)
    assert [t["so_hv"] for t in tk] == [5, 5, 0]
    assert "sheet 2" in tk[0]["dung_chung"] and "sheet 1" in tk[1]["dung_chung"] and tk[2]["dung_chung"] == ""
    wb = openpyxl.load_workbook(io.BytesIO(ket_qua))
    ws = wb["1"]
    assert ws["C14"].value == "18000" and ws["C18"].value == "18004"
    assert ws["B18"].value is not None and ws["B19"].value is None  # STT dừng ở học viên thứ 5
    assert "Trưởng Trung tâm" in str(ws["C19"].value)  # chân trang dời lên liền sau dòng cuối
    assert wb["2"]["C14"].value == "18000"
    ml = wb["Muc luc"]
    assert ml["F2"].value == "Số HV" and ml["F3"].value == "=COUNTA('1'!C14:C18)"
    assert ml["F5"].value.startswith("=COUNTA('3'!C14:")  # sheet chưa điền: đếm trên bảng mẫu → 0
    assert ml["F6"].value == "=COUNTA('4'!C14:C43)"  # sheet dự phòng: bảng 30 dòng, không co, vẫn đếm Số HV
    assert wb["4"]["B43"].value is not None and not any(t["sheet"] == "4" for t in tk)


def test_gom_file_khung():
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("a.xlsx", b"1")
        z.writestr("__MACOSX/a.xlsx", b"x")
        z.writestr("~$tmp.xlsx", b"x")
        z.writestr("sub/b.xlsx", b"2")
        z.writestr("c.txt", b"x")
    r = khdt_hv.gom_file_khung([("a.xlsx", b"OLD"), ("z.xlsx", b"Z")], [("p.zip", zb.getvalue()), ("d.xlsx", b"4")])
    assert sorted(r) == [("a.xlsx", b"1"), ("b.xlsx", b"2"), ("d.xlsx", b"4"), ("z.xlsx", b"Z")]


if __name__ == "__main__":
    test_doc_dshv()
    test_dien_khung_nhanh_gian_bang()
    test_thu_gon_bang_so_hv_dung_chung()
    test_gom_file_khung()
    print("OK")
